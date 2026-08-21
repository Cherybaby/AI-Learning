"""从模型 Excel 生成 RTE 读取声明与故障处理逻辑。

脚本读取 model/SF020_SysFaultMngt.xlsx 中 Rport 工作表的 A、B 两列（跳过标题行），
按行轮循并更新 src/SF020_SysFaultMngt.c 里的 SF020_SysFaultMngt_Cycle_5ms：

轮循规则：
1) A 列数字相对上一行发生变化时，才处理该行（生成对应的 Merged_Code）。
2) A 列数字未变化（或 A 为空/合并单元格续行）时，跳过该行，不生成任何代码。
3) 对于被处理的行，若 B 列为 Reserved / TBD（或空），则跳过其变量声明/读取/触发/恢复逻辑，
   但仍需生成 Monr_SF020_u8_SysFaultIndex_UL++，以保持故障索引与槽位对齐。
4) 其余有效信号：生成 UINT8 <name>; 声明、Rte_Read 调用，以及触发/恢复/else 逻辑块 + index++。
"""
from pathlib import Path
import re
from openpyxl import load_workbook

# 根据脚本所在目录定位到 [SF020_SysFaultMngt.xlsx](http://_vscodecontentref_/0)
script_dir = Path(__file__).resolve().parent
xlsx_path = script_dir.parent / "model" / "SF020_SysFaultMngt.xlsx"

wb = load_workbook(xlsx_path, data_only=True, read_only=True)
ws = wb["Rport"]

# 读取 A 列（分组编号）与 B 列（信号名），跳过标题行
rows_data = []
for row in ws.iter_rows(min_row=2, min_col=1, max_col=2, values_only=True):
    rows_data.append((row[0], row[1]))
wb.close()


def _norm_a(v):
    """归一化 A 列取值用于比较：整数浮点(1.0)转成整数字符串('1')，空白转 None。"""
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return s if s else None


SKIP_TOKENS = {"RESERVED", "TBD"}

# 轮循 A/B 列：
#   - A 列数字相对上一行“变化” -> 处理该行
#   - A 列数字“未变化”，或 A 为空/合并单元格续行 -> 跳过该行（不生成代码）
#   - 被处理的行里，B 为 Reserved/TBD/空 -> 跳过其声明/读取/逻辑，但仍需 index++
processed_rows = []          # 元素: {"kind": "signal"|"skip", "name": str}
_prev_a_key = None
for a_val, b_val in rows_data:
    a_key = _norm_a(a_val)
    if a_key is None or a_key == _prev_a_key:
        continue             # 空白/续行/与上一行相同 -> 整行跳过
    _prev_a_key = a_key
    b_str = "" if b_val is None else str(b_val).strip()
    if b_str == "" or b_str.upper() in SKIP_TOKENS:
        processed_rows.append({"kind": "skip", "name": b_str or "Reserved"})
    else:
        processed_rows.append({"kind": "signal", "name": b_str})

signal_names = [r["name"] for r in processed_rows if r["kind"] == "signal"]

invalid_names = [n for n in signal_names if not re.match(r"^[A-Za-z_][A-Za-z0-9_]*$", n)]
if invalid_names:
    raise ValueError(f"发现不合法的 C 标识符: {invalid_names}")

print("A/B 原始数据行数：", len(rows_data))
print("A 变化后参与处理的行：", [(r["kind"], r["name"]) for r in processed_rows])
print("有效信号名：", signal_names)

# 顶部固定代码
ResetIndex_Code = ["  Monr_SF020_u8_SysFaultIndex_UL = 0U;"]
ReturnValue_Code = ["  boolean value;\n"
                    "  boolean* ret = &value;"]

# 变量声明与读取：仅针对有效信号（Reserved/TBD 不声明、不读取）
VariableDeclare_Code = [f"  UINT8 {name};" for name in signal_names]
ReadInput_Code = [f"  (void)Rte_Read_{name}_gdu8(&{name});" for name in signal_names]


def _fault_trigger(name):
    return (
        f"  if (({name} >= 1U) && (Monr_SF020_b_SysFaultStatus_UL[Monr_SF020_u8_SysFaultIndex_UL] == 0U))\n"
        "  {\n"
        f"    TriggerFault(Monr_SF020_u8_SysFaultIndex_UL, {name}, ret);\n"
        "  }"
    )


def _fault_recover(name):
    return (
        f"  else if (({name} == 0U) && (Monr_SF020_b_SysFaultStatus_UL[Monr_SF020_u8_SysFaultIndex_UL] == 1U))\n"
        "  {\n"
        "    RecoverFault(Monr_SF020_u8_SysFaultIndex_UL, ret);\n"
        "  }"
    )


_ELSE_BLOCK = ("  else\n"
               "  {\n"
               "    /* no action */\n"
               "  }")
_INDEX_UPDATE = "  Monr_SF020_u8_SysFaultIndex_UL++;"
BlockStart = "  /* CODEGEN_MERGED_CODE_START */"
BlockEnd = "  /* CODEGEN_MERGED_CODE_END */"


# 按要求顺序拼接代码：
# ResetIndex_Code -> ReturnValue_Code -> VariableDeclare_Code -> ReadInput_Code ->
# 逐行：signal -> 触发/恢复/else + index++；skip(Reserved/TBD) -> 仅 index++
Merged_Code = []
Merged_Code.extend(ResetIndex_Code)
Merged_Code.extend(ReturnValue_Code)
Merged_Code.append("")
Merged_Code.extend(VariableDeclare_Code)
Merged_Code.append("")
Merged_Code.extend(ReadInput_Code)
Merged_Code.append("")

for idx_val, r in enumerate(processed_rows):
    if r["kind"] == "signal":
        name = r["name"]
        Merged_Code.append(f"  /* Monr_SF020_u8_SysFaultIndex_UL = {idx_val} : {name} */")
        Merged_Code.append(_fault_trigger(name))
        Merged_Code.append(_fault_recover(name))
        Merged_Code.append(_ELSE_BLOCK)
        Merged_Code.append(_INDEX_UPDATE)
        Merged_Code.append("")
    else:
        # Reserved/TBD：跳过读取与判断逻辑，但索引仍需前进
        Merged_Code.append(
            f"  /* Monr_SF020_u8_SysFaultIndex_UL = {idx_val} : {r['name']} "
            "(reserved/TBD, skip logic, advance index only) */"
        )
        Merged_Code.append(_INDEX_UPDATE)
        Merged_Code.append("")


def insert_code_at_function_start(c_text: str, func_signature: str, new_body_lines: list[str]) -> str:
    signature_index = c_text.find(func_signature)
    if signature_index == -1:
        raise ValueError(f"未找到函数签名: {func_signature}")

    brace_start = c_text.find("{", signature_index)
    if brace_start == -1:
        raise ValueError(f"未找到函数起始大括号: {func_signature}")

    newline = "\r\n" if "\r\n" in c_text else "\n"
    block_lines = [BlockStart, *new_body_lines, BlockEnd, ""]
    new_block = newline + newline.join(block_lines)

    block_start_index = c_text.find(BlockStart, brace_start)
    block_end_index = c_text.find(BlockEnd, brace_start)
    if block_start_index != -1 and block_end_index != -1:
        block_end_index += len(BlockEnd)
        if block_end_index < len(c_text) and c_text[block_end_index:block_end_index + len(newline)] == newline:
            block_end_index += len(newline)
        return c_text[:block_start_index] + new_block + c_text[block_end_index:]

    return c_text[: brace_start + 1] + new_block + c_text[brace_start + 1 :]


c_file_path = script_dir.parent / "src" / "SF020_SysFaultMngt.c"
h_file_path = script_dir.parent / "include" / "SF020_SysFaultMngt.h"
c_text = c_file_path.read_text(encoding="utf-8")
updated_c_text = insert_code_at_function_start(
    c_text,
    "void SF020_SysFaultMngt_Cycle_5ms(void)",
    Merged_Code,
)
updated_c_text = updated_c_text.replace(
    "Rte_Call_DiagnosticMonitor_SetEventStatus",
    "Dem_SetEventStatus",
)
updated_c_text = updated_c_text.replace(
    "DEM_EVENT_STATUS_FAILED",
    "Monr_SF020_u8_DTCIndex_UL, DEM_EVENT_STATUS_FAILED",
)
updated_c_text = updated_c_text.replace(
    "DEM_EVENT_STATUS_PASSED",
    "Monr_SF020_u8_DTCIndex_UL, DEM_EVENT_STATUS_PASSED",
)
updated_c_text = updated_c_text.replace(
    "static void SF020_SysFaultMngt_NoXTCRecover(boolean *rty_Ret);",
    "static void SF020_SysFaultMngt_NoXTCRecover(boolean *rty_Ret);\nvoid RecoverFault(uint8 XtcId, boolean *Ret);\nvoid TriggerFault(uint8 XtcId, uint8 Para, boolean *Ret);",
)
c_file_path.write_text(updated_c_text, encoding="utf-8")
print(f"\n已更新: {c_file_path}")

h_text = h_file_path.read_text(encoding="utf-8")
updated_h_text = h_text.replace(
    '#include "XPSbW_PublicCal.h"',
    '#include "XPSbW_PublicCal.h"\n#include "Dem.h"',
)
h_file_path.write_text(updated_h_text, encoding="utf-8")
print(f"\n已更新: {h_file_path}")


print("\n拼接后的完整代码：")
for code in Merged_Code:
    print(code)